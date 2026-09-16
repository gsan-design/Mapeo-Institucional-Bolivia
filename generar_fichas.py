#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera las fichas de programas (HTML estatico para GitHub Pages)
a partir del Excel de mapeo.

Uso:
    python generar_fichas.py SAT_prueba.xlsx
    python generar_fichas.py SAT_prueba.xlsx --hoja "Mapeo v3" --salida index.html

Estructura que espera del Excel (una fila de encabezados en la fila 1):
    A  Nombre del programa o instrumento      -> combinada/vacia por todo el bloque del programa
    B  Periodo de implementacion
    C  Descripcion
    D  Categoria principal de intervencion
    F  Problema (generico) que busca resolver -> combinada/vacia por bloque de problema
    G  Hipotesis (generica)                   -> combinada/vacia por bloque de problema
    H  Instrumentos o herramientas            -> UNA FILA POR INSTRUMENTO
       formato sugerido: "Nombre del instrumento: descripcion..."

Reglas:
  - Celda combinada  -> el valor se propaga a todas sus filas.
  - Celda vacia      -> hereda el valor de la fila de arriba (mismo programa).
  - Fila sin nombre de programa heredable -> se ignora.
"""

import argparse
import html
import json
import re
from pathlib import Path

from openpyxl import load_workbook

# Encabezado exacto en el Excel -> clave interna
COLUMNAS = {
    "Nombre del programa o instrumento": "programa",
    "Período de implementación": "periodo",
    "Descripción": "descripcion",
    "Categoría principal de intervención": "categoria",
    "Categoría(s) secundaria(s) de intervención": "categorias_sec",
    "Problema (genérico) que busca resolver": "problema",
    "Hipótesis (genérica)": "hipotesis",
    "Instrumentos o herramientas": "instrumento",
}

HEREDABLES = ["programa", "periodo", "descripcion", "categoria",
              "categorias_sec", "problema", "hipotesis"]


def limpiar(v):
    if v is None:
        return ""
    s = str(v).replace("\\n", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return "" if s.lower() in ("nan", "none") else s


def leer_filas(ruta, hoja=None):
    """Devuelve la tabla como lista de dicts, con combinadas y vacias resueltas."""
    wb = load_workbook(ruta, data_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]

    # 1. Propagar el valor de cada celda combinada a todo su rango.
    valores = {}
    for fila in ws.iter_rows():
        for c in fila:
            valores[(c.row, c.column)] = c.value
    for rango in ws.merged_cells.ranges:
        ancla = valores.get((rango.min_row, rango.min_col))
        for r in range(rango.min_row, rango.max_row + 1):
            for c in range(rango.min_col, rango.max_col + 1):
                valores[(r, c)] = ancla

    # 2. Mapear encabezados -> indice de columna.
    encabezados = {}
    for c in range(1, ws.max_column + 1):
        titulo = limpiar(valores.get((1, c)))
        if titulo in COLUMNAS:
            encabezados[COLUMNAS[titulo]] = c
    faltan = set(COLUMNAS.values()) - set(encabezados)
    if "programa" in faltan or "instrumento" in faltan:
        raise SystemExit(f"Faltan columnas obligatorias en la hoja: {faltan}")

    # 3. Leer filas heredando hacia abajo.
    filas, anterior = [], {}
    for r in range(2, ws.max_row + 1):
        actual = {k: limpiar(valores.get((r, c))) for k, c in encabezados.items()}
        if not any(actual.values()):
            continue
        # Un nombre de programa nuevo corta la herencia.
        if actual.get("programa") and actual["programa"] != anterior.get("programa"):
            anterior = {}
        for k in HEREDABLES:
            if not actual.get(k):
                actual[k] = anterior.get(k, "")
        if not actual.get("programa"):
            continue
        filas.append(actual)
        anterior = actual
    return filas


def agrupar(filas):
    """filas planas -> [programa -> problemas -> instrumentos]"""
    programas, idx_prog = [], {}
    for f in filas:
        clave = f["programa"]
        if clave not in idx_prog:
            idx_prog[clave] = len(programas)
            programas.append({
                "nombre": f["programa"],
                "periodo": f["periodo"],
                "descripcion": f["descripcion"],
                "categoria": f["categoria"],
                "categorias_sec": f.get("categorias_sec", ""),
                "problemas": [],
            })
        prog = programas[idx_prog[clave]]

        if not f["problema"]:
            continue
        prob = next((p for p in prog["problemas"] if p["texto"] == f["problema"]), None)
        if prob is None:
            prob = {"texto": f["problema"], "hipotesis": f["hipotesis"], "instrumentos": []}
            prog["problemas"].append(prob)

        if f["instrumento"]:
            nombre, _, desc = f["instrumento"].partition(":")
            prob["instrumentos"].append({
                "nombre": nombre.strip(" .") or f["instrumento"],
                "descripcion": desc.strip(" .") if desc.strip() else "",
            })
    return programas


PLANTILLA = """<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>__TITULO__</title>
<style>
  :root{
    --fondo:#ffffff; --tinta:#111111; --borde:#111111;
    --chip:#ffffff; --chip-hover:#f1f1f1; --tenue:#555555; --panel:#fafafa;
  }
  :root:not([data-theme="light"]){ }
  @media (prefers-color-scheme: dark){
    :root:not([data-theme="light"]){
      --fondo:#141414; --tinta:#f2f2f2; --borde:#f2f2f2;
      --chip:#1c1c1c; --chip-hover:#2a2a2a; --tenue:#b5b5b5; --panel:#1a1a1a;
    }
  }
  :root[data-theme="dark"]{
    --fondo:#141414; --tinta:#f2f2f2; --borde:#f2f2f2;
    --chip:#1c1c1c; --chip-hover:#2a2a2a; --tenue:#b5b5b5; --panel:#1a1a1a;
  }
  *{box-sizing:border-box}
  body{
    margin:0; padding:24px 16px; background:var(--fondo); color:var(--tinta);
    font-family:"Segoe UI",Calibri,system-ui,-apple-system,"Helvetica Neue",Arial,sans-serif;
    line-height:1.45;
  }
  .contenedor{max-width:1100px; margin:0 auto; display:flex; flex-direction:column; gap:28px}
  .ficha{
    border:3px solid var(--borde); border-radius:22px; padding:26px 28px 30px;
    background:var(--fondo);
  }
  .ficha h2{margin:0 0 18px; font-size:1.3rem; font-weight:700}
  .campo{margin-bottom:14px; max-width:70ch}
  .campo .etiqueta{font-weight:700; display:block}
  .campo .valor{font-weight:600}
  .seccion{font-weight:700; margin:22px 0 12px}
  .rejilla{display:grid; grid-template-columns:repeat(auto-fit,minmax(240px,1fr)); gap:16px}
  .chip{
    border:3px solid var(--borde); border-radius:14px; background:var(--chip);
    color:inherit; font:inherit; font-weight:700; text-align:left;
    padding:16px 18px; cursor:pointer; min-height:78px;
    display:flex; align-items:center; transition:background .15s;
  }
  .chip:hover,.chip:focus-visible{background:var(--chip-hover)}
  .detalle{border:3px solid var(--borde); border-radius:18px; padding:22px 24px; background:var(--panel)}
  .detalle .bloque{margin-bottom:22px; max-width:70ch}
  .detalle .bloque:last-child{margin-bottom:0}
  .instrumentos{display:flex; flex-wrap:wrap; gap:12px; margin-top:12px}
  .chip-inst{
    border:2px solid var(--borde); border-radius:12px; background:var(--chip);
    color:inherit; font:inherit; font-weight:700; font-size:.9rem; text-align:left;
    padding:12px 14px; cursor:pointer; max-width:280px; transition:background .15s;
  }
  .chip-inst:hover{background:var(--chip-hover)}
  .desc-inst{
    margin-top:14px; padding:14px 16px; border-left:4px solid var(--borde);
    background:var(--fondo); font-size:.92rem; max-width:70ch;
  }
  .volver{
    margin-bottom:16px; border:2px solid var(--borde); border-radius:999px;
    background:transparent; color:inherit; font:inherit; font-weight:700;
    font-size:.85rem; padding:7px 16px; cursor:pointer;
  }
  .volver:hover{background:var(--chip-hover)}
  .oculto{display:none}
  @media print{ .chip{page-break-inside:avoid} }
</style>
</head>
<body>
<div class="contenedor" id="app"></div>

<script>
const PROGRAMAS = __DATOS__;

const app = document.getElementById("app");

function esc(t){ const d=document.createElement("div"); d.textContent=t||""; return d.innerHTML; }

function render(){
  app.innerHTML = PROGRAMAS.map((p,i) => `
    <section class="ficha" id="ficha-${i}">
      <h2>${esc(p.nombre)}${p.periodo ? " (" + esc(p.periodo) + ")" : ""}</h2>
      <div id="cuerpo-${i}"></div>
    </section>`).join("");
  PROGRAMAS.forEach((_,i) => vistaResumen(i));
}

function vistaResumen(i){
  const p = PROGRAMAS[i];
  const cuerpo = document.getElementById("cuerpo-" + i);
  cuerpo.innerHTML = `
    ${p.descripcion ? `<div class="campo"><span class="etiqueta">Descripción breve:</span><span class="valor">${esc(p.descripcion)}</span></div>` : ""}
    ${p.categoria ? `<div class="campo"><span class="etiqueta">Categoría principal de intervención:</span><span class="valor">${esc(p.categoria)}</span></div>` : ""}
    <div class="seccion">Problemas que busca resolver</div>
    <div class="rejilla">
      ${p.problemas.map((pr,j) => `
        <button class="chip" data-prog="${i}" data-prob="${j}">
          Problema ${j+1}. ${esc(titulito(pr.texto))}
        </button>`).join("")}
    </div>`;
  cuerpo.querySelectorAll(".chip").forEach(b =>
    b.addEventListener("click", () => vistaDetalle(i, +b.dataset.prob)));
}

function titulito(t){
  const corte = t.indexOf(":");
  return corte > 0 && corte < 90 ? t.slice(0, corte) : t;
}

function vistaDetalle(i, j){
  const p = PROGRAMAS[i], pr = p.problemas[j];
  const cuerpo = document.getElementById("cuerpo-" + i);
  cuerpo.innerHTML = `
    <button class="volver">← Volver</button>
    <div class="detalle">
      <div class="bloque"><strong>Problema ${j+1}.</strong> ${esc(pr.texto)}</div>
      ${pr.hipotesis ? `<div class="bloque"><strong>Hipótesis de solución:</strong> ${esc(pr.hipotesis)}</div>` : ""}
      <div class="bloque">
        <strong>Instrumentos y herramientas:</strong>
        <div class="instrumentos">
          ${pr.instrumentos.map((it,k) => `<button class="chip-inst" data-i="${k}">${esc(it.nombre)}</button>`).join("")}
        </div>
        <div class="desc-inst oculto" id="desc-inst"></div>
      </div>
    </div>`;
  cuerpo.querySelector(".volver").addEventListener("click", () => vistaResumen(i));
  const caja = cuerpo.querySelector("#desc-inst");
  cuerpo.querySelectorAll(".chip-inst").forEach(b => b.addEventListener("click", () => {
    const it = pr.instrumentos[+b.dataset.i];
    caja.innerHTML = `<strong>${esc(it.nombre)}.</strong> ${esc(it.descripcion || "Sin descripción registrada.")}`;
    caja.classList.remove("oculto");
  }));
  document.getElementById("ficha-" + i).scrollIntoView({behavior:"smooth", block:"start"});
}

render();
</script>
</body>
</html>
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel", help="ruta al archivo .xlsx")
    ap.add_argument("--hoja", default=None, help='nombre de la hoja (por defecto la primera)')
    ap.add_argument("--salida", default="index.html")
    ap.add_argument("--titulo", default="Fichas de programas")
    args = ap.parse_args()

    programas = agrupar(leer_filas(args.excel, args.hoja))
    datos = json.dumps(programas, ensure_ascii=False, indent=1)

    salida = Path(args.salida)
    salida.write_text(
        PLANTILLA.replace("__DATOS__", datos).replace("__TITULO__", html.escape(args.titulo)),
        encoding="utf-8",
    )
    Path(salida.with_name("datos.json")).write_text(datos, encoding="utf-8")

    print(f"OK -> {salida} ({len(programas)} programa(s))")
    for p in programas:
        print(f"  · {p['nombre']}: {len(p['problemas'])} problemas, "
              f"{sum(len(x['instrumentos']) for x in p['problemas'])} instrumentos")


if __name__ == "__main__":
    main()
